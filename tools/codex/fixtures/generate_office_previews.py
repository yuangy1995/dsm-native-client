#!/usr/bin/env python3
"""生成供原生 Quick Look 合成测试使用的 Office 文件；不含真实资料。

使用工作区已有的 python-docx、openpyxl 与 python-pptx，不改变应用依赖。
"""
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from openpyxl import Workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[3]
DESTINATION = ROOT / "apple/Apps/DsmMac/Tests/Fixtures/Office"
CREATED = datetime(2026, 10, 2, 0, 0, 0)


def normalize(source: Path, destination: Path) -> None:
    """固定 ZIP 时间与文件顺序，避免每次生成产生无关二进制差异。"""
    with ZipFile(source) as original, ZipFile(destination, "w") as output:
        for name in sorted(original.namelist()):
            info = ZipInfo(name, (2026, 10, 2, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            output.writestr(info, original.read(name))


def generate() -> None:
    DESTINATION.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix="office-preview-fixtures-") as directory:
        temporary = Path(directory)
        word = Document()
        word.add_heading("Office preview — 合成文档", 0)
        word.add_paragraph("Synthetic Word document. 本页仅用于预览测试。")
        word.core_properties.author = "Synthetic fixture"
        word.core_properties.last_modified_by = "Synthetic fixture"
        word.core_properties.created = CREATED
        word.core_properties.modified = CREATED
        word.save(temporary / "sample.docx")

        excel = Workbook()
        sheet = excel.active
        sheet.title = "Synthetic"
        sheet.append(["Item / 项目", "Amount / 数量"])
        sheet.append(["Preview", 12])
        sheet.append(["合成数据", 34])
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 22
        excel.properties.creator = "Synthetic fixture"
        excel.properties.lastModifiedBy = "Synthetic fixture"
        excel.properties.created = CREATED
        excel.properties.modified = CREATED
        excel.save(temporary / "sample.xlsx")

        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[0])
        slide.shapes.title.text = "Office preview"
        slide.placeholders[1].text = "Synthetic PowerPoint\n合成演示文稿"
        slides.core_properties.author = "Synthetic fixture"
        slides.core_properties.last_modified_by = "Synthetic fixture"
        slides.core_properties.created = CREATED
        slides.core_properties.modified = CREATED
        slides.save(temporary / "sample.pptx")
        for extension in ["docx", "xlsx", "pptx"]:
            name = "sample." + extension
            normalize(temporary / name, DESTINATION / name)


if __name__ == "__main__":
    generate()
